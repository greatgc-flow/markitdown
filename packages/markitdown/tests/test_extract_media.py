"""Synthetic original-media extraction contracts."""

import hashlib
import io
import re
from pathlib import Path

import pytest
from PIL import Image
from markitdown import MarkItDown, StreamInfo
from markitdown.converters import (
    DocxConverter,
    PptxConverter,
    XlsxConverter,
    HwpConverter,
)
from ._hwpx_builder import build_hwpx


def png():
    out = io.BytesIO()
    Image.new("RGB", (200, 100), "red").save(out, "PNG")
    return out.getvalue()


def source(fmt):
    out = io.BytesIO()
    if fmt == "docx":
        from docx import Document

        doc = Document()
        for _ in range(2):
            doc.add_picture(io.BytesIO(png()))
        doc.save(out)
    elif fmt == "pptx":
        from pptx import Presentation

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        for x in (0, 100):
            slide.shapes.add_picture(io.BytesIO(png()), x, 0)
        doc.save(out)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage

        doc = Workbook()
        doc.active.append(["synthetic"])
        for anchor in ("A2", "C2"):
            doc.active.add_image(XLImage(io.BytesIO(png())), anchor)
        doc.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    "fmt,converter",
    [("docx", DocxConverter), ("pptx", PptxConverter), ("xlsx", XlsxConverter)],
)
def test_original_bytes_and_duplicate_occurrences(tmp_path, fmt, converter):
    data = source(fmt)
    assets = tmp_path / "assets"
    result = converter().convert(
        io.BytesIO(data),
        StreamInfo(extension="." + fmt),
        extract_media=assets,
        media_base_dir=tmp_path,
    )
    expected = hashlib.sha256(png()).hexdigest()[:12] + ".png"
    assert [p.name for p in assets.iterdir()] == [expected]
    assert (assets / expected).read_bytes() == png()
    assert result.markdown.count("assets/" + expected) == 2
    assert "data:" not in result.markdown


def test_extractor_rejects_collision_and_unsafe_extension(tmp_path):
    from markitdown._media import save_media

    name = hashlib.sha256(png()).hexdigest()[:12] + ".png"
    (tmp_path / name).write_bytes(b"collision")
    with pytest.raises(ValueError, match="collision"):
        save_media(png(), "png", extract_media=tmp_path)
    with pytest.raises(ValueError):
        save_media(png(), "../outside", extract_media=tmp_path)


def test_hwpx_skipped_wmf_is_explicit_and_bytes_preserved(tmp_path):
    import zipfile
    from .test_ocr_triage_media import package

    with zipfile.ZipFile(io.BytesIO(build_hwpx())) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries["BinData/image1.png"] = png()
    entries["BinData/skipped.wmf"] = b"synthetic unsupported WMF"
    result = HwpConverter().convert(
        io.BytesIO(package(entries)),
        StreamInfo(extension=".hwpx"),
        extract_media=tmp_path,
    )
    assert "[IMAGE" not in result.markdown
    assert "![WMF omitted]" in result.markdown
    assert any(p.read_bytes() == png() for p in tmp_path.iterdir())


def test_python_options_exclusive(tmp_path):
    with pytest.raises(ValueError):
        DocxConverter().convert(
            io.BytesIO(source("docx")),
            StreamInfo(extension=".docx"),
            extract_media=tmp_path,
            keep_data_uris=True,
        )


def test_dedupe_identical_bytes_even_with_extension_alias(tmp_path):
    from markitdown._media import save_media

    a = save_media(png(), "jpg", extract_media=tmp_path)
    b = save_media(png(), "jpeg", extract_media=tmp_path)
    assert a == b
    assert len(list(tmp_path.iterdir())) == 1


@pytest.mark.parametrize("member", ["word/media/../../escape.png", "word/media/A.png"])
def test_unsafe_package_members_fail_without_zip_fallback(tmp_path, member):
    import zipfile
    from .test_ocr_triage_media import package
    from markitdown import FileConversionException

    with zipfile.ZipFile(io.BytesIO(source("docx"))) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries["word/media/a.png"] = png()
    entries[member] = png()
    with pytest.raises(FileConversionException):
        MarkItDown().convert_stream(
            io.BytesIO(package(entries)),
            stream_info=StreamInfo(extension=".docx"),
            extract_media=tmp_path / "assets",
        )


def test_cli_links_relative_to_output_file(tmp_path):
    import subprocess, sys

    path = tmp_path / "synthetic.docx"
    path.write_bytes(source("docx"))
    output = tmp_path / "nested" / "out.md"
    output.parent.mkdir()
    assets = tmp_path / "assets"
    done = subprocess.run(
        [
            sys.executable,
            "-m",
            "markitdown",
            str(path),
            "-o",
            str(output),
            "--extract-media",
            str(assets),
        ],
        capture_output=True,
    )
    assert done.returncode == 0, done.stderr
    assert done.stdout == b""
    assert "../assets/" in output.read_text(encoding="utf-8")


def test_hwp5_payload_extraction_and_repeated_named_markers(monkeypatch, tmp_path):
    import hwp_hwpx_parser as parser
    from types import SimpleNamespace

    closed = []
    reader = SimpleNamespace(
        get_images=lambda: [
            SimpleNamespace(filename="BIN0001.png", data=png(), format="png", index=0)
        ],
        extract_text_with_notes=lambda options: SimpleNamespace(
            text="[IMAGE: BIN0001.png]\n[IMAGE: BIN0001.png]", footnotes=[], endnotes=[]
        ),
        close=lambda: closed.append(True),
    )
    monkeypatch.setattr(parser, "read", lambda path: reader)
    result = HwpConverter().convert(
        io.BytesIO(b"synthetic HWP5"),
        StreamInfo(extension=".hwp"),
        extract_media=tmp_path,
    )
    assert result.markdown.count(hashlib.sha256(png()).hexdigest()[:12]) == 2
    assert len(list(tmp_path.iterdir())) == 1
    assert next(tmp_path.iterdir()).read_bytes() == png()
    assert closed == [True]


def test_missing_docx_image_is_explicit(tmp_path):
    import zipfile
    from .test_ocr_triage_media import package

    with zipfile.ZipFile(io.BytesIO(source("docx"))) as z:
        entries = {n: z.read(n) for n in z.namelist() if n != "word/media/image1.png"}
    result = DocxConverter().convert(
        io.BytesIO(package(entries)),
        StreamInfo(extension=".docx"),
        extract_media=tmp_path,
    )
    assert "![Image omitted]" in result.markdown
    assert result.unsupported_media is True


def test_real_synthetic_hwp5_parser_and_extraction(tmp_path):
    import hwp_hwpx_parser as parser
    from ._hwp5_builder import build_hwp5
    from markitdown.ocr_triage import triage

    data = build_hwp5(png())
    path = tmp_path / "synthetic.hwp"
    path.write_bytes(data)
    reader = parser.read(path)
    try:
        images = reader.get_images()
        assert len(images) == 1  # WMF is actually skipped by parser 1.0.x.
        assert (
            images[0].data == png()
        )  # Stored deflate bytes are not the image payload.
    finally:
        reader.close()
    report = triage(io.BytesIO(data))
    assert report["file"]["format"] == "hwp"
    assert [r["id"] for r in report["media"]] == [
        "BinData/BIN0001.png",
        "BinData/BIN0002.wmf",
    ]
    assert report["media"][0]["payload_sha256"] == hashlib.sha256(png()).hexdigest()
    assert report["media"][0]["stored_bytes"] == 4096
    assert report["media"][0]["stored_sha256"] != report["media"][0]["payload_sha256"]
    assert any(i["code"] == "unsupported_media" for i in report["issues"])
    assets = tmp_path / "assets"
    result = HwpConverter().convert(
        io.BytesIO(data), StreamInfo(extension=".hwp"), extract_media=assets
    )
    assert next(assets.iterdir()).read_bytes() == png()
    assert "![WMF omitted]" in result.markdown


def test_python_mutual_exclusion_applies_before_converter_selection(tmp_path):
    with pytest.raises(ValueError):
        MarkItDown().convert_stream(
            io.BytesIO(b"text"),
            stream_info=StreamInfo(extension=".txt"),
            extract_media=tmp_path,
            keep_data_uris=True,
        )


def test_hwp_reference_alias_collision_is_rejected(tmp_path):
    from markitdown._hwp_media import extract_hwp_markdown

    rows = [
        ("BinData/same.png", png(), "png", True),
        ("BinData/same.jpg", png() + b"different", "jpg", True),
    ]
    with pytest.raises(ValueError, match="colliding"):
        extract_hwp_markdown("[IMAGE: same]", rows, extract_media=tmp_path)
