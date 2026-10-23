"""Phase 1b: synthetic package media, safety, and deterministic evidence."""

import io
import json
import os
import socket
import struct
import subprocess
import sys
import zipfile

import pytest
from PIL import Image

from markitdown import ocr_triage as core
from ._hwpx_builder import build_hwpx

R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/package/2006/relationships"


def png(size=(200, 100)):
    out = io.BytesIO()
    Image.new("RGB", size).save(out, "PNG")
    return out.getvalue()


def package(entries):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, value in entries.items():
            z.writestr(name, value)
    return out.getvalue()


def rels(*targets):
    return (
        f'<Relationships xmlns="{P}">'
        + "".join(
            f'<Relationship Id="r{i}" Type="{R}/image" Target="{target}" {mode}/>'
            for i, (target, mode) in enumerate(targets)
        )
        + "</Relationships>"
    )


def refs(*ids):
    return (
        f'<root xmlns:r="{R}">'
        + "".join(f'<image r:embed="r{i}"/>' for i in ids)
        + "</root>"
    )


def docx():
    return {
        "word/document.xml": refs(0, 0, 1, 2, 3).replace(
            "</root>",
            '<w:t xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            + "text " * 30
            + "</w:t></root>",
        ),
        "word/_rels/document.xml.rels": rels(
            ("media/a.png", ""),
            ("media/icon.png", ""),
            ("media/banner.png", ""),
            ("media/a.emf", ""),
        ),
        "word/media/a.png": png(),
        "word/media/icon.png": png((16, 16)),
        "word/media/banner.png": png((600, 30)),
        "word/media/a.emf": b"unknown vector",
        "word/header1.xml": refs(0, 1),
        "word/_rels/header1.xml.rels": rels(
            ("media/header.png", ""), ("media/a.png", "")
        ),
        "word/media/header.png": png(),
    }


def test_docx_occurrences_thresholds_and_decor():
    data = package(docx())
    report = core.triage(io.BytesIO(data))
    assert report["file"]["format"] == "docx"
    assert report["decision"] == "review"
    assert report["coverage"]["state"] == "complete"
    rows = report["media"]
    assert len(rows) == 7
    assert sum(r["id"] == "word/media/a.png" for r in rows) == 3
    assert all(r["where"] is None for r in rows)
    assert [r["id"] for r in rows if r["decor"]] == ["word/media/header.png"]
    assert report["summary"]["media_small"] == 1
    assert report["summary"]["image_occurrences"] == 7
    assert len(report["candidates"]) == 6
    assert all(
        c["reasons"]
        == (["decor_media"] if c["id"].endswith("header.png") else ["embedded_image"])
        and c["priority"] == "low"
        for c in report["candidates"]
    )
    assert next(r for r in rows if r["id"].endswith("emf"))["px"] is None
    assert list(rows[0]) == ["id", "where", "px", "bytes", "ext", "decor", "references"]


def test_pptx_presentation_order_and_layout_master():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png()), 0, 0, width=Inches(1))
        slide.shapes.add_picture(io.BytesIO(png()), 0, 0, width=Inches(1))
    out = io.BytesIO()
    prs.save(out)
    with zipfile.ZipFile(out) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    # Reverse presentation order, keeping slide filenames unchanged.
    text = entries["ppt/presentation.xml"].decode()
    import re

    slides = re.findall(r"<p:sldId\b[^>]*/>", text)
    assert len(slides) == 2
    entries["ppt/presentation.xml"] = text.replace(
        "".join(slides), "".join(reversed(slides))
    )
    for part in ("slideMasters/slideMaster1", "slideLayouts/slideLayout1"):
        entries[f"ppt/{part}.xml"] = refs(0)
        folder, name = part.split("/")
        entries[f"ppt/{folder}/_rels/{name}.xml.rels"] = rels(
            ("../media/decor.png", "")
        )
    entries["ppt/media/decor.png"] = png()
    report = core.triage(io.BytesIO(package(entries)))
    assert report["file"]["format"] == "pptx"
    placed = [r for r in report["media"] if not r["decor"]]
    assert len(placed) == 4
    assert [(r["where"], r["references"][0]["part"]) for r in placed] == [
        ({"slide": 1}, "ppt/slides/slide2.xml"),
        ({"slide": 1}, "ppt/slides/slide2.xml"),
        ({"slide": 2}, "ppt/slides/slide1.xml"),
        ({"slide": 2}, "ppt/slides/slide1.xml"),
    ]
    assert sum(r["decor"] for r in report["media"]) == 2
    assert report["decision"] == "ocr_recommended"


def test_xlsx_drawing_and_hwpx_builder():
    entries = {
        "xl/workbook.xml": "<root/>",
        "xl/drawings/drawing1.xml": refs(0, 0),
        "xl/drawings/_rels/drawing1.xml.rels": rels(("../media/a.png", "")),
        "xl/media/a.png": png(),
    }
    report = core.triage(io.BytesIO(package(entries)))
    assert report["file"]["format"] == "xlsx"
    assert len(report["media"]) == 2
    assert all(r["where"] is None for r in report["media"])
    with zipfile.ZipFile(io.BytesIO(build_hwpx())) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    assert b'binaryItemIDRef="image1"' in entries["Contents/section0.xml"]
    entries["BinData/image1.png"] = png()
    report = core.triage(io.BytesIO(package(entries)))
    assert report["file"]["format"] == "hwpx"
    assert report["media"][0]["id"] == "BinData/image1.png"
    assert report["media"][0]["references"] == [
        {"part": "Contents/section0.xml", "id": "image1"}
    ]


def test_external_corrupt_ole_and_missing_pillow(monkeypatch):
    entries = docx()
    entries["word/_rels/document.xml.rels"] = rels(
        ("https://invalid/image.png", 'TargetMode="External"')
    )
    monkeypatch.setattr(socket, "socket", lambda *a, **k: pytest.fail("network"))
    report = core.triage(io.BytesIO(package(entries)))
    assert "external_media" in [i["code"] for i in report["issues"]]
    assert report["coverage"]["state"] == "partial"
    for data, code in [
        (b"PK\x03\x04broken", "corrupt"),
        (b"\xd0\xcf\x11\xe0" + b"x" * 20, "unsupported_format"),
    ]:
        assert core.triage(io.BytesIO(data))["issues"][0]["code"] == code
    data = package(docx())
    monkeypatch.setitem(sys.modules, "PIL", None)
    report = core.triage(io.BytesIO(data))
    assert len(report["candidates"]) == 7
    assert all(r["px"] is None for r in report["media"])


@pytest.mark.parametrize(
    "limit,value",
    [
        ("MAX_ZIP_ENTRIES", 3),
        ("MAX_ZIP_TOTAL_BYTES", 100),
        ("MAX_ZIP_READ_BYTES", 100),
        ("MAX_XML_BYTES", 100),
        ("MAX_MEDIA_REFERENCES", 1),
    ],
)
def test_package_caps(monkeypatch, limit, value):
    monkeypatch.setattr(core, limit, value)
    report = core.triage(io.BytesIO(package(docx())))
    assert report["coverage"]["state"] != "complete"
    assert any(i["code"] == "zip_limit" for i in report["issues"])


def test_bomb_and_many_entries(monkeypatch):
    entries = docx()
    entries["word/document.xml"] = b"0" * (5 * 1024**2)
    monkeypatch.setattr(core, "MAX_XML_BYTES", 1024)
    reads = []
    original = zipfile.ZipExtFile.read

    def read(self, n=-1):
        assert 0 <= n <= 1024
        data = original(self, n)
        reads.append(len(data))
        return data

    monkeypatch.setattr(zipfile.ZipExtFile, "read", read)
    report = core.triage(io.BytesIO(package(entries)))
    assert report["coverage"]["state"] == "partial"
    assert sum(reads) <= core.MAX_ZIP_READ_BYTES
    data = package({f"empty/{i}": b"" for i in range(100000)})
    monkeypatch.setattr(core, "MAX_ZIP_ENTRIES", 100)
    report = core.triage(io.BytesIO(data))
    assert report["decision"] == "unknown"
    assert report["issues"][0]["code"] == "zip_limit"


def test_media_header_only_and_restore(monkeypatch):
    entries = docx()
    entries["word/media/a.png"] += b"0" * (5 * 1024**2)
    monkeypatch.setattr(core, "MAX_MEDIA_PROBE_BYTES", 64)
    stream = io.BytesIO(package(entries))
    stream.seek(17)
    report = core.triage(stream)
    assert stream.tell() == 17 and not stream.closed
    assert report["decision"] == "review"
    assert next(r for r in report["media"] if r["id"].endswith("a.png"))["px"] == [
        200,
        100,
    ]


def test_media_deterministic_subprocess(tmp_path):
    path = tmp_path / "synthetic.docx"
    path.write_bytes(package(docx()))
    script = "import json,sys;from markitdown.ocr_triage import triage;print(json.dumps(triage(sys.argv[1]),ensure_ascii=False,separators=(',',':')))"
    outputs = [
        subprocess.check_output(
            [sys.executable, "-c", script, str(path)],
            env={**os.environ, "PYTHONHASHSEED": seed},
        )
        for seed in ("1", "987")
    ]
    assert outputs[0] == outputs[1]
    assert json.loads(outputs[0])["file"]["format"] == "docx"


def test_corrupt_member_reads_still_spend_budget(monkeypatch):
    data = bytearray(
        package(
            {
                "word/document.xml": "<root/>",
                **{f"word/media/{i}.png": png() + b"0" * 1500 for i in range(10)},
            }
        )
    )
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        offset = z.start_dir
        for item in z.infolist():
            if "/media/" in item.filename:
                # Wrong CRC forces failure after real decompression of every image.
                data[offset + 16] ^= 1
            lengths = struct.unpack_from("<3H", data, offset + 28)
            offset += 46 + sum(lengths)
    monkeypatch.setattr(core, "MAX_MEDIA_READ_BYTES", 4096, raising=False)
    read_sizes = []
    original = zipfile.ZipExtFile.read

    def watched(self, n=-1):
        read_sizes.append(n)
        return original(self, n)

    monkeypatch.setattr(zipfile.ZipExtFile, "read", watched)
    report = core.triage(io.BytesIO(data))
    assert report["coverage"]["state"] == "partial"
    assert sum(read_sizes) <= 4096 + len(b"<root/>")
    assert len(report["candidates"]) == 10
    assert all(row["px"] is None for row in report["media"])


@pytest.mark.parametrize(
    "size,candidate",
    [((24, 625), True), ((24, 624), False), ((23, 1000), False), ((600, 30), True)],
)
def test_embedded_exact_size_thresholds(size, candidate):
    report = core.triage(
        io.BytesIO(
            package({"word/document.xml": "<root/>", "word/media/a.png": png(size)})
        )
    )
    assert bool(report["candidates"]) == candidate
    assert report["summary"]["media_small"] == int(not candidate)


def test_unreadable_xml_entities_and_mixed_decor_are_conservative():
    entries = docx()
    entries["word/document.xml"] = (
        '<!DOCTYPE root [<!ENTITY x "private">]><root>&x;</root>'
    )
    report = core.triage(io.BytesIO(package(entries)))
    assert report["coverage"]["state"] == "partial"
    assert not any(row["decor"] for row in report["media"])
    assert "corrupt" in [i["code"] for i in report["issues"]]


def test_hwpx_manifest_alias_and_occurrences():
    with zipfile.ZipFile(io.BytesIO(build_hwpx())) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries["Contents/content.hpf"] = entries["Contents/content.hpf"].replace(
        b"<opf:manifest>",
        b'<opf:manifest><opf:item id="image1" href="../BinData/picture.png"/>',
    )
    entries["Contents/section0.xml"] = entries["Contents/section0.xml"].replace(
        b'<hc:img binaryItemIDRef="image1"/>', b'<hc:img binaryItemIDRef="image1"/>' * 2
    )
    entries["BinData/picture.png"] = png()
    report = core.triage(io.BytesIO(package(entries)))
    assert [r["id"] for r in report["media"]] == ["BinData/picture.png"] * 2
    assert len(report["candidates"]) == 2


def test_media_never_uses_converter_or_pixel_decode(monkeypatch):
    import markitdown

    data = package(docx())

    def forbidden(*args, **kwargs):
        pytest.fail("converter or pixel decoding")

    monkeypatch.setattr(markitdown, "MarkItDown", forbidden)
    monkeypatch.setattr(Image.Image, "load", forbidden)
    assert core.triage(io.BytesIO(data))["decision"] == "review"


def test_real_decompression_obeys_probe_budget(monkeypatch):
    entries = {
        "word/document.xml": "<root/>",
        "word/media/a.png": png() + b"0" * (5 * 1024**2),
    }
    data = package(entries)
    monkeypatch.setattr(core, "MAX_MEDIA_PROBE_BYTES", 64)
    original = zipfile._get_decompressor
    output_bytes = []

    class Watched:
        def __init__(self, inner):
            self.inner = inner

        def __getattr__(self, key):
            return getattr(self.inner, key)

        def decompress(self, *args):
            result = self.inner.decompress(*args)
            output_bytes.append(len(result))
            return result

    monkeypatch.setattr(
        zipfile, "_get_decompressor", lambda *args: Watched(original(*args))
    )
    report = core.triage(io.BytesIO(data))
    assert report["media"][0]["px"] == [200, 100]
    assert sum(output_bytes) <= 64 + len(b"<root/>")


def test_unbounded_zip_compression_method_is_not_read(monkeypatch):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_LZMA) as z:
        z.writestr("word/document.xml", "<root/>")
        z.writestr("word/media/a.png", png() + b"0" * (5 * 1024**2))

    def watched(self, *args, **kwargs):
        pytest.fail("unbounded compression method opened")

    monkeypatch.setattr(zipfile.ZipFile, "open", watched)
    report = core.triage(io.BytesIO(out.getvalue()))
    assert report["coverage"]["state"] == "partial"
    assert any(i["code"] == "unsupported_compression" for i in report["issues"])
    assert len(report["candidates"]) == 1


def test_directory_offset_cannot_redirect_preflight(monkeypatch):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        z.writestr("word/document.xml", "<root/>")
        z.writestr("padding", b"0" * 1000)
        z.writestr("word/media/a.png", png())
    data = bytearray(out.getvalue())
    eocd = data.rfind(b"PK\x05\x06")
    length = struct.unpack_from("<I", data, eocd + 12)[0]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        member = z.getinfo("padding")
    fake = member.header_offset + 30 + len(member.filename)
    header = bytearray(46)
    header[:4] = b"PK\x01\x02"
    struct.pack_into("<H", header, 28, length - 46)
    data[fake : fake + 46] = header
    struct.pack_into("<HH", data, eocd + 8, 1, 1)
    struct.pack_into("<I", data, eocd + 16, fake)
    monkeypatch.setattr(core, "MAX_ZIP_ENTRIES", 2)

    def forbidden(*args, **kwargs):
        pytest.fail("directory preflight was redirected before ZipFile allocation")

    monkeypatch.setattr(zipfile, "ZipFile", forbidden)
    report = core.triage(io.BytesIO(data))
    assert report["decision"] == "unknown"
    assert report["issues"][0]["code"] == "corrupt"


def test_non_image_media_is_excluded_but_unknown_image_type_is_kept():
    entries = {
        "word/document.xml": refs(0, 1),
        "word/_rels/document.xml.rels": rels(
            ("media/picture.bin", ""), ("media/movie.mp4", "")
        ).replace(
            f'{R}/image" Target="media/movie.mp4', f'{R}/video" Target="media/movie.mp4'
        ),
        "word/media/picture.bin": b"unknown image encoding",
        "word/media/movie.mp4": b"video",
        "word/media/audio.mp3": b"audio",
    }
    report = core.triage(io.BytesIO(package(entries)))
    assert [r["id"] for r in report["media"]] == ["word/media/picture.bin"]
    assert report["candidates"][0]["id"] == "word/media/picture.bin"
    with zipfile.ZipFile(io.BytesIO(build_hwpx())) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries["BinData/image1.bin"] = b"unknown image encoding"
    entries["BinData/sound.mp3"] = b"audio"
    report = core.triage(io.BytesIO(package(entries)))
    assert [r["id"] for r in report["media"]] == ["BinData/image1.bin"]


@pytest.mark.parametrize("fmt", ["docx", "pptx", "hwpx", "xlsx"])
@pytest.mark.parametrize("chars", [0, 29, 30, 60])
def test_office_scan_like_text_threshold(fmt, chars):
    root, marker, text_ns = {
        "docx": (
            "word/",
            "document.xml",
            "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        ),
        "pptx": (
            "ppt/",
            "presentation.xml",
            "http://schemas.openxmlformats.org/drawingml/2006/main",
        ),
        "hwpx": (
            "Contents/",
            "content.hpf",
            "http://www.hancom.co.kr/hwpml/2011/paragraph",
        ),
        "xlsx": (
            "xl/",
            "workbook.xml",
            "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        ),
    }[fmt]
    part = {
        "docx": "document.xml",
        "pptx": "slides/slide1.xml",
        "hwpx": "section0.xml",
        "xlsx": "sharedStrings.xml",
    }[fmt]
    text = f'<root xmlns:t="{text_ns}"><t:t>' + " x \t\n" * chars + "</t:t></root>"
    entries = {
        root + marker: "<root/>",
        root + part: text,
        ("BinData/" if fmt == "hwpx" else root + "media/") + "scan.png": png(),
    }
    report = core.triage(io.BytesIO(package(entries)))
    assert report["coverage"]["state"] == "complete"
    assert report["summary"]["nonspace_chars"] == chars
    assert report["candidates"][0]["priority"] == ("high" if chars < 30 else "low")
    assert report["decision"] == ("ocr_recommended" if chars < 30 else "review")


def test_many_images_have_separate_budget_and_keep_decor():
    entries = {
        "word/document.xml": "<root/>",
        "word/header1.xml": refs(*range(300)),
        "word/_rels/header1.xml.rels": rels(
            *[(f"media/{i}.png", "") for i in range(300)]
        ),
    }
    entries.update(
        {f"word/media/{i}.png": png() + b"0" * (64 * 1024) for i in range(300)}
    )
    report = core.triage(io.BytesIO(package(entries)))
    assert report["coverage"]["state"] == "complete"
    assert not report["issues"]
    assert len(report["candidates"]) == 300
    assert all(row["decor"] for row in report["media"])
    assert all(
        c["priority"] == "low" and c["reasons"] == ["decor_media"] and c["decor"]
        for c in report["candidates"]
    )


def test_media_budget_exhaustion_is_partial(monkeypatch):
    monkeypatch.setattr(core, "MAX_MEDIA_READ_BYTES", 100, raising=False)
    report = core.triage(
        io.BytesIO(package({"word/document.xml": "<root/>", "word/media/a.png": png()}))
    )
    assert report["coverage"]["limits"]["max_media_read_bytes"] == 100
    assert report["coverage"]["state"] == "partial"
    assert any(i["code"] == "zip_limit" for i in report["issues"])


def test_unresolved_embed_is_partial():
    report = core.triage(io.BytesIO(package({"word/document.xml": refs(99)})))
    assert report["coverage"]["state"] == "partial"
    assert report["decision"] == "unknown"
    assert any(i["code"] == "unresolved_embed" for i in report["issues"])


def test_hwpx_wmf_is_unsupported_issue():
    report = core.triage(
        io.BytesIO(
            package(
                {
                    "Contents/content.hpf": "<root/>",
                    "BinData/skipped.wmf": b"synthetic WMF",
                }
            )
        )
    )
    assert report["media"][0]["id"] == "BinData/skipped.wmf"
    assert any(i["code"] == "unsupported_media" for i in report["issues"])


def test_hwp5_parser_images_stable_and_not_decor(monkeypatch):
    from types import SimpleNamespace
    import hwp_hwpx_parser as parser

    reader = SimpleNamespace(
        get_images=lambda: [
            SimpleNamespace(filename="BIN0001.png", data=png(), format="png", index=0),
            SimpleNamespace(filename="BIN0002.wmf", data=b"wmf", format="wmf", index=1),
        ],
        close=lambda: None,
    )
    monkeypatch.setattr(parser, "read", lambda path: reader)
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"synthetic HWP5"
    stream = io.BytesIO(data)
    stream.seek(3)
    report = core.triage(stream)
    assert stream.tell() == 3
    assert report["file"]["format"] == "hwp"
    assert [r["id"] for r in report["media"]] == [
        "BinData/BIN0001.png",
        "BinData/BIN0002.wmf",
    ]
    assert all(r["decor"] is False for r in report["media"])
    assert all(c["decor"] is False for c in report["candidates"])
    assert report["summary"]["image_occurrences"] == 2
    assert any(i["code"] == "unsupported_media" for i in report["issues"])
    assert core.triage(io.BytesIO(data)) == report
